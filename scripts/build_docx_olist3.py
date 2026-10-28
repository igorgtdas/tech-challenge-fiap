"""Gera o relatório ABNT do OLIST-3 (Word) a partir de reports/achados_olist3.json.

Uso (na raiz do projeto, com o .venv ativo e depois de rodar o notebook 02):
    python scripts/build_docx_olist3.py

Usa o mesmo template FIAP e as mesmas funções de formatação do relatório do OLIST-2 (build_docx.py).
"""
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

import build_docx as b
from olist import config

OUT = b.ABNT / "OLIST-3_relatorio.docx"
DIMENSOES_ROTULO = {"núcleo": "Núcleo", "diferencial": "Diferencial"}


def build():
    d = json.loads((config.REPORTS / "achados_olist3.json").read_text(encoding="utf-8"))
    achados, chave, metricas = d["achados"], d["chave"], {m["id"]: m for m in d["metricas"]}
    o2 = json.loads(config.JSON_PATH.read_text(encoding="utf-8"))
    I = {i["id"]: i for i in o2["indicadores"]}
    A = {x["id"]: x for x in achados}
    n_dim = sum(1 for c in d["cobertura"] if c["coberta"] == "sim")
    fig = lambda name: config.FIGURES / f"{name}.png"  # noqa: E731

    doc = Document(str(b.make_template()))
    body_el = doc.element.body
    for el in list(body_el):
        if el.tag != qn("w:sectPr"):
            body_el.remove(el)
    cp = doc.core_properties
    cp.author, cp.title = "Diogo, Igor e Lorran", "OLIST-3 · Achados de negócio no dataset Olist"
    cp.comments, cp.last_modified_by = "Tech Challenge Agentic AI – Fase 1 – FIAP", "Igor"

    # Título e autores
    p = b._para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.5, after=6)
    b._font(p.add_run("Onde estão os problemas e as oportunidades da Olist: achados de negócio no dataset para orientar o uso de agentes de IA"), 12, bold=True)
    b.blank(doc)
    b.rich(b._para(doc, align=WD_ALIGN_PARAGRAPH.CENTER), "Diogo^1^; Igor^1*^; Lorran^1^", base_bold=True)
    b.blank(doc)
    b.rich(b._para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.0), "^1^ FIAP. Pós-graduação em Agentic AI. Tech Challenge – Fase 1. São Paulo, SP, Brasil", size=9)
    b.rich(b._para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.0), "*autor correspondente: igorgtdas1@gmail.com · card OLIST-3", size=9)
    b.blank(doc)

    # Sumário executivo
    k = {c["rank"]: c for c in chave}
    b.heading(doc, "Sumário Executivo")
    b.body(doc, (
        f"Este documento transforma os indicadores-base do dataset Olist em achados de negócio. Foram feitas {len(achados)} análises: seis núcleo, "
        f"sugeridas pelos desafios do enunciado, e seis diferenciais, que mostram onde e quando o problema nasce. Elas cobrem {n_dim} das 8 dimensões do dataset. "
        f"O achado mais forte é que o atraso aparece junto com a insatisfação: {k[1]['resumo']} "
        f"O atraso dá sinal antes de acontecer, porque se forma no transporte e na postagem fora do prazo pelo vendedor. Os comentários negativos, "
        f"com milhares de textos, falam principalmente de entrega e de produto. Geografia, distância e picos de poucos meses concentram o problema, "
        f"e a recompra é baixa mesmo entre clientes satisfeitos. Todos os achados são associações, não relações de causa, e cada número tem fórmula "
        f"do Excel que o reproduz. Os seis achados-chave orientam a matriz de oportunidades, os KPIs e o mapa de agentes."
    ))
    b.rich(b._para(doc), "**Palavras-chave:** achados de negócio; logística de e-commerce; satisfação do cliente; análise de avaliações; agentes de IA.")
    b.blank(doc)

    p = b._para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    b._font(p.add_run("Where Olist's problems and opportunities lie: business findings from the dataset to guide AI agents"), 11, bold=True)
    p = b._para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    b._font(p.add_run("Abstract"), 11, bold=True)
    b.body(doc, (
        f"This document turns the Olist baseline indicators into business findings through {len(achados)} analyses covering {n_dim} of the 8 dataset dimensions. "
        "Late deliveries are strongly associated with low review scores, delays build up mostly in transport and in late shipping by sellers, negative comments "
        "focus on delivery and product issues, and delays concentrate in distant regions and in a few peak months. Repurchase is low regardless of satisfaction. "
        "All findings are associations, not causal claims, and every figure is reproducible with Excel formulas."
    ))
    b.rich(b._para(doc), "**Keywords:** business insights; e-commerce logistics; customer satisfaction; review analysis; AI agents.")
    b.blank(doc)

    # Introdução
    b.heading(doc, "Introdução")
    b.body(doc, (
        "O enunciado do Tech Challenge pede que o relatório apresente os principais problemas ou oportunidades da Olist e as análises feitas no dataset. "
        "O grupo é avaliado pela profundidade das análises e pela qualidade dos insights (FIAP, 2026). A etapa anterior, o card OLIST-2, preparou a base "
        f"e calculou indicadores como a taxa de atraso ({I['I05']['texto']}) e a nota média ({I['I07']['texto']})."
    ))
    b.body(doc, (
        "Indicadores dizem quanto; achados dizem onde, quando e com quem o problema acontece, e o que isso muda para o negócio. Este trabalho corresponde "
        "ao card OLIST-3. Seu objetivo é transformar os números em achados que justifiquem a escolha dos agentes de IA nos próximos cards: a matriz de "
        "oportunidades (OLIST-5), os KPIs (OLIST-6) e o mapa de agentes (OLIST-8)."
    ))

    # Metodologia
    b.heading(doc, "Metodologia")
    b.body(doc, (
        "A análise é quantitativa, exploratória e descritiva, feita sobre a base com uma linha por pedido do OLIST-2, em que as oito regras de qualidade "
        "já estão aplicadas. A essa base foram acrescentadas, ainda por pedido, informações sobre:"
    ))
    b.bullet(doc, "o vendedor e a categoria do item principal;")
    b.bullet(doc, "o peso dos itens e a participação do frete no valor;")
    b.bullet(doc, "o tempo de cada etapa do ciclo do pedido;")
    b.bullet(doc, "se a postagem respeitou o prazo-limite do vendedor;")
    b.bullet(doc, "a distância em linha reta entre os CEPs de vendedor e cliente;")
    b.bullet(doc, "a forma de pagamento e a recompra.")
    b.body(doc, (
        "Os comentários das avaliações com nota 1 ou 2 foram classificados em seis temas por palavras-chave, definidas a partir da leitura de uma "
        "amostra. É uma classificação aproximada, usada para dimensionar os temas. A Tabela 1 lista as análises."
    ))
    b.caption(doc, "Tabela 1. Análises realizadas, tipo e dimensões do dataset cobertas.")
    b.table(doc, ["ID", "Tipo", "Pergunta", "Dimensões"],
            [[x["id"], DIMENSOES_ROTULO[x["tipo"]], x["pergunta"], x["dimensoes"]] for x in achados], [1.2, 2.2, 7.6, 5.0], size=8)
    b.source(doc)
    b.body(doc, (
        "Cada análise é apresentada em quatro partes:"
    ))
    b.bullet(doc, "**Fato:** o que os números mostram;")
    b.bullet(doc, "**Leitura:** o que o fato sugere sobre o negócio;")
    b.bullet(doc, "**Cautela:** até onde o dado permite concluir;")
    b.bullet(doc, "**E daí?:** o que muda para a Olist.")
    b.body(doc, (
        "Como os dados são observacionais, nenhuma conclusão afirma causa: o texto usa expressões como \"associado a\" e \"compatível com\", e uma checagem "
        f"automática procura linguagem causal. Cada número citado vem de uma métrica com fórmula do Excel na planilha de achados. As fórmulas das {len(metricas)} métricas foram "
        "recalculadas fora do Python e chegaram aos mesmos valores."
    ))

    # Resultados
    b.heading(doc, "Resultados e Discussão")
    figs = {"N1": "Percentual de notas 1 ou 2 por faixa de atraso.", "N3": "Temas dos comentários negativos.",
            "N6": "Taxa de atraso por região do cliente.", "D1": "Tempo médio por etapa do ciclo do pedido.",
            "D2": "Distribuição da taxa de atraso dos vendedores com 50 ou mais entregas.",
            "D3": "Taxa de atraso por peso do pedido e por categoria.", "D5": "Pedidos entregues e taxa de atraso por mês.",
            "D6": "Taxa de atraso e prazo por distância entre vendedor e cliente."}
    n_fig = 0
    for grupo, titulo in (("núcleo", "Análises núcleo"), ("diferencial", "Análises diferenciais")):
        b.heading(doc, titulo, 2)
        for x in [a for a in achados if a["tipo"] == grupo]:
            p = b._para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=3, keep=True)
            b._font(p.add_run(f"{x['id']}. {x['titulo']}"), 11, bold=True)
            for rot, campo in (("Fato", "fato"), ("Leitura", "leitura"), ("Cautela", "cautela"), ("E daí?", "e_dai")):
                b.rich(b._para(doc, indent=True), f"**{rot}.** {x[campo]}")
            if x["figura"]:
                n_fig += 1
                b.figure(doc, n_fig, figs[x["id"]], fig(x["figura"]), 14.5)

    b.heading(doc, "Achados-chave", 2)
    b.body(doc, (
        "A Tabela 2 reúne os seis achados que mais pesam na escolha dos agentes. Eles estão em ordem de relevância: primeiro os que afetam mais "
        "clientes e têm sinal mais forte nos dados. Cada um indica o desafio do enunciado que ilustra e o tipo de solução que pede. "
        "A escolha dos agentes em si fica para o mapa de agentes."
    ))
    b.caption(doc, "Tabela 2. Achados-chave, desafios do enunciado e oportunidade associada.")
    b.table(doc, ["#", "Achado", "Evidência", "Desafios do enunciado", "Oportunidade"],
            [[c["rank"], c["titulo"], c["resumo"], c["desafios"], c["oportunidade"]] for c in chave], [0.7, 3.0, 5.3, 3.2, 3.8], size=8)
    b.source(doc)

    # Considerações finais
    b.heading(doc, "Considerações Finais")
    b.body(doc, (
        "Os achados apontam para um mesmo eixo: a entrega. O atraso aparece junto com a maior parte da insatisfação, se forma em etapas que podem ser "
        "acompanhadas e se concentra em rotas, vendedores e meses identificáveis. Isso favorece agentes que monitoram e agem antes do problema, e não só "
        "depois da reclamação. Os comentários em texto e a recompra baixa abrem uma segunda frente, de atendimento e de retenção, em que a IA generativa "
        "tem papel natural."
    ))
    b.body(doc, (
        "Os dados vão de 2016 a 2018 e não trazem transportadora, custos, motivos de cancelamento nem registros de atendimento. Por isso, os achados devem "
        "ser lidos como hipóteses fortes a validar em piloto, e não como relações de causa."
    ))

    b.heading(doc, "Referências")
    refs = [
        "FIAP. Tech Challenge Agentic AI – Fase 1: Olist Intelligent Marketplace. São Paulo: FIAP, 2026. Enunciado da atividade.",
        "MCKINNEY, W. Data structures for statistical computing in Python. In: PYTHON IN SCIENCE CONFERENCE, 9., 2010, Austin. Proceedings [...]. Austin: SciPy, 2010. p. 56-61.",
        "OLIST. Brazilian E-Commerce Public Dataset by Olist. Kaggle, 2018. Disponível em: https://www.kaggle.com/datasets/olistbr/brazilian-ecommerce. Acesso em: 26 out. 2026.",
        "TUKEY, J. W. Exploratory data analysis. Reading: Addison-Wesley, 1977.",
    ]
    for ref in refs:
        p = b._para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.0, after=12)
        b._font(p.add_run(ref), 11)

    doc.save(OUT)
    print("Word:", OUT)
    return OUT


if __name__ == "__main__":
    build()
