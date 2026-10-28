"""Gera o Word "Insights e análises de Oportunidades" com o mesmo conteúdo e a mesma ordem da página do OLIST-3 no Confluence.

Uso (na raiz do projeto, com o .venv ativo e depois de rodar o notebook 02):
    python scripts/build_docx_insights.py

Usa o template FIAP e as funções de formatação de build_docx.py. Os textos vêm de reports/achados_olist3.json,
a mesma fonte da página do Confluence.
"""
import json

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

import build_docx as b
from olist import config

OUT = b.ABNT / "OLIST-3_insights_oportunidades.docx"
FIG_CAP = {
    "o3_n1_atraso_nota": "% de notas 1 ou 2 por faixa de atraso",
    "o3_n3_temas_comentarios": "Temas dos comentários negativos",
    "o3_n6_atraso_regiao": "Taxa de atraso por região",
    "o3_d1_etapas_ciclo": "Tempo médio por etapa do ciclo do pedido",
    "o3_d2_vendedores": "Taxa de atraso dos vendedores com 50+ entregas",
    "o3_d3_peso_categoria": "Atraso por peso e por categoria",
    "o3_d5_sazonalidade": "Pedidos e taxa de atraso por mês",
    "o3_d6_distancia": "Atraso e prazo por distância vendedor–cliente",
}
ARQUIVOS = [
    ("notebooks/02_olist3_achados.ipynb", "as 12 análises passo a passo, os achados e as checagens"),
    ("reports/achados_olist3.xlsx", "Leia-me, Indicadores (métricas), Achados-chave, Achados, base_pedidos ampliada e uma aba por métrica (M01–M12)"),
    ("reports/achados_olist3.json", "achados e métricas para Word e Confluence"),
    ("reports/figures/o3_*.png", "os 8 gráficos"),
    ("docs/abnt/OLIST-3_relatorio.docx", "relatório no padrão ABNT (template FIAP)"),
    ("src/olist/findings.py", "registro dos achados e checagem de linguagem causal"),
]


def build():
    d = json.loads((config.REPORTS / "achados_olist3.json").read_text(encoding="utf-8"))
    achados, chave, metricas = d["achados"], d["chave"], d["metricas"]
    n_dim = sum(1 for c in d["cobertura"] if c["coberta"] == "sim")
    n_dif = sum(1 for a in achados if a["tipo"] == "diferencial")

    doc = Document(str(b.make_template()))
    body_el = doc.element.body
    for el in list(body_el):
        if el.tag != qn("w:sectPr"):
            body_el.remove(el)
    cp = doc.core_properties
    cp.author, cp.title = "Diogo, Igor e Lorran", "Insights e análises de Oportunidades"
    cp.comments, cp.last_modified_by = "Tech Challenge Agentic AI – Fase 1 – FIAP · card OLIST-3", "Igor"

    # Título e autores
    p = b._para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.5, after=6)
    b._font(p.add_run("Insights e análises de Oportunidades"), 14, bold=True)
    b.rich(b._para(doc, align=WD_ALIGN_PARAGRAPH.CENTER), "Diogo; Igor; Lorran", base_bold=True)
    b.rich(b._para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0, after=6), "Tech Challenge Agentic AI – Fase 1 · FIAP · card OLIST-3", size=9)

    # Introdução
    b.heading(doc, "Introdução")
    b.body(doc, (
        f"Este documento reúne os insights do dataset público da Olist e as oportunidades que eles apontam para o uso de IA e de agentes. "
        f"Ele é a etapa que transforma os indicadores-base em achados de negócio. Foram feitas {len(achados)} análises sobre os dados de pedidos, "
        f"entregas, avaliações, vendedores, produtos, pagamentos e geolocalização: {len(achados) - n_dif} análises núcleo, ligadas aos desafios do enunciado, "
        f"e {n_dif} diferenciais, que mostram onde e quando o problema nasce."
    ))
    b.body(doc, (
        "O documento abre com os seis achados-chave, em ordem de relevância. Depois vem cada análise no formato Fato, Leitura, Cautela e E daí?. "
        "Esses achados são a base da matriz de oportunidades, dos KPIs e do mapa de agentes. Todos os números saem dos dados e podem ser conferidos na planilha de achados."
    ))
    b.rich(b._para(doc, indent=True), (
        "**O que o enunciado pede:** os \"principais problemas ou oportunidades identificadas\" e as \"análises realizadas no dataset\" (Enunciado, p. 7). "
        "Isso é avaliado no critério 4, profundidade das análises e qualidade dos insights (p. 5)."
    ))

    # Achados-chave
    b.heading(doc, "Achados-chave (para Diogo e Lorran)")
    b.body(doc, "Os 6 achados que mais pesam na escolha dos agentes, em ordem de relevância. A coluna \"oportunidade\" diz que tipo de solução o achado pede; a escolha do agente fica para o OLIST-8.")
    b.caption(doc, "Tabela 1. Achados-chave.")
    b.table(doc, ["#", "Achado", "Evidência", "Desafios do enunciado", "Quem usa", "Oportunidade"],
            [[c["rank"], c["titulo"], c["resumo"], c["desafios"], c["para_quem"], c["oportunidade"]] for c in chave],
            [0.6, 2.4, 3.6, 2.6, 3.4, 3.4], size=7.5)
    b.source(doc)

    # Como ler
    b.heading(doc, "Como ler cada achado")
    b.caption(doc, "Tabela 2. As quatro partes de cada achado.")
    b.table(doc, ["Parte", "Pergunta que responde"], [
        ["Fato", "O que os números mostram, sem opinião?"],
        ["Leitura", "O que isso sugere sobre o negócio?"],
        ["Cautela", "Até onde o dado permite concluir? (associação não é causa)"],
        ["E daí?", "O que muda para a Olist e para os agentes?"],
    ], [3.0, 13.0])
    b.source(doc)
    b.body(doc, (
        f"Todo número citado vem de uma métrica (M01–M{len(metricas):02d}) com fórmula do Excel na planilha de achados. "
        "As fórmulas foram recalculadas fora do Python e bateram. Uma checagem automática garante que nenhum texto afirma causa."
    ))

    # Análises
    n_tab, n_fig = 2, 0
    for tipo, titulo in (("núcleo", "Análises núcleo (N1–N6)"), ("diferencial", "Análises diferenciais (D1–D6)")):
        b.heading(doc, titulo)
        if tipo == "diferencial":
            b.body(doc, "Análises que vão além do óbvio e não estão no relatório de referência: mostram onde e quando o problema nasce.")
        for a in [x for x in achados if x["tipo"] == tipo]:
            b.heading(doc, f"{a['id']} · {a['titulo']}", 2)
            p = b._para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.0, after=4)
            for rot, val in (("Pergunta: ", a["pergunta"]), (" · Dimensões: ", a["dimensoes"]),
                             (" · Desafios: ", a["desafios"]), (" · Métricas: ", a["metricas"])):
                b._font(p.add_run(rot), 9, italic=True)
                b._font(p.add_run(val), 9)
            n_tab += 1
            b.caption(doc, f"Tabela {n_tab}. {a['id']}: fato, leitura, cautela e e daí.")
            b.table(doc, ["Parte", "Conteúdo"], [["Fato", a["fato"]], ["Leitura", a["leitura"]],
                                                ["Cautela", a["cautela"]], ["E daí?", a["e_dai"]]], [2.2, 13.8], size=9)
            b.source(doc, "Fonte: Elaboração própria a partir do dataset Olist (Kaggle).")
            if a["figura"]:
                n_fig += 1
                b.figure(doc, n_fig, FIG_CAP[a["figura"]] + ".", config.FIGURES / f"{a['figura']}.png", 14.5)

    # Critérios de aceite
    b.heading(doc, "Critérios de aceite do OLIST-3")
    itens = [
        (True, f"CA1 · Análises no dataset cobrindo ≥ 5 das 8 dimensões do enunciado ({n_dim} de 8)"),
        (True, f"CA2 · ≥ 6 análises, das quais ≥ 2 diferenciais ({len(achados)} análises, {n_dif} diferenciais)"),
        (True, "CA3 · Todas com Fato, Leitura, Cautela e E daí? preenchidos"),
        (True, "CA4 · Nenhuma frase afirma causalidade (checagem automática)"),
        (True, "CA5 · Cada achado-chave aponta o desafio do enunciado que ilustra"),
        (True, f"CA6 · Todo número tem fórmula na planilha ({len(metricas)} métricas em reports/achados_olist3.xlsx)"),
        (False, "CA7 · Lista dos achados-chave entregue ao Diogo (gráficos e KPIs) e ao Lorran (oportunidades e agentes)"),
    ]
    for ok, texto in itens:
        p = b._para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.15, after=2)
        p.paragraph_format.left_indent = b.Cm(0.6)
        b._font(p.add_run(("☑ " if ok else "☐ ")), 11)
        b._font(p.add_run(texto), 11)

    # Arquivos
    b.heading(doc, "Onde estão os arquivos")
    n_tab += 1
    b.caption(doc, f"Tabela {n_tab}. Arquivos do card no repositório do projeto.")
    b.table(doc, ["Arquivo", "O que é"], [list(x) for x in ARQUIVOS], [6.0, 10.0], size=9)
    b.source(doc)

    doc.save(OUT)
    print("Word:", OUT)
    return OUT


if __name__ == "__main__":
    build()
